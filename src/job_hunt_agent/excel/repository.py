from __future__ import annotations

import copy
import hashlib
import os
import re
import shutil
import threading
import time
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import BinaryIO, Callable, Final, Iterator

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet

from job_hunt_agent.excel.models import (
    TrackerApplication,
    TrackerApplicationDraft,
    TrackerApplicationPatch,
)


HEADERS: Final = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)
FIELDS: Final = (
    "company",
    "role",
    "applied_date",
    "location",
    "status",
    "next_step",
    "next_time",
    "job_url",
    "notes",
)
CLOSED_MARKERS: Final = ("未通过", "结束", "拒绝", "放弃", "终止", "淘汰")
DIVIDER_COLOR: Final = "E7E6E3"
LOCK_TIMEOUT_SECONDS = 10.0
_LOCK_POLL_SECONDS: Final = 0.05
_THREAD_LOCKS: dict[str, threading.Lock] = {}
_THREAD_LOCKS_GUARD = threading.Lock()


@dataclass(frozen=True)
class _LocatedApplication:
    row: int
    record: TrackerApplication


@dataclass
class _CellSnapshot:
    value: object
    style: object
    hyperlink: object
    comment: object


@dataclass
class _RowSnapshot:
    source_row: int | None
    cells: tuple[_CellSnapshot, ...]
    height: float | None
    hidden: bool
    outline_level: int
    collapsed: bool


@dataclass(frozen=True)
class _MergeSnapshot:
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    anchor: _CellSnapshot


def _normalize(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", value)
    return " ".join(normalized.split()).casefold()


def _application_id(company: str, role: str, applied_date: date | None) -> str:
    normalized_date = applied_date.isoformat() if applied_date is not None else ""
    payload = "\x1f".join((_normalize(company), _normalize(role), normalized_date))
    return f"app_{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:16]}"


def _as_date(value: object) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(str(value).strip())


def _is_closed(status: object) -> bool:
    normalized = str(status or "").strip()
    return bool(normalized) and any(marker in normalized for marker in CLOSED_MARKERS)


def _cell_snapshot(cell: Cell | MergedCell) -> _CellSnapshot:
    return _CellSnapshot(
        value=cell.value,
        style=copy.copy(cell._style),
        hyperlink=copy.copy(getattr(cell, "hyperlink", None)),
        comment=copy.copy(getattr(cell, "comment", None)),
    )


def _write_cell_snapshot(cell: Cell, snapshot: _CellSnapshot) -> None:
    cell.value = snapshot.value
    cell._style = copy.copy(snapshot.style)
    cell.hyperlink = copy.copy(snapshot.hyperlink)
    cell.comment = copy.copy(snapshot.comment)


def _try_lock_file(handle: BinaryIO) -> bool:
    handle.seek(0)
    try:
        if os.name == "nt":
            import msvcrt

            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        return False
    return True


def _unlock_file(handle: BinaryIO) -> None:
    handle.seek(0)
    if os.name == "nt":
        import msvcrt

        msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
    else:
        import fcntl

        fcntl.flock(handle.fileno(), fcntl.LOCK_UN)


@contextmanager
def _exclusive_workbook_lock(workbook_path: Path) -> Iterator[None]:
    lock_key = os.path.normcase(str(workbook_path.resolve()))
    with _THREAD_LOCKS_GUARD:
        thread_lock = _THREAD_LOCKS.setdefault(lock_key, threading.Lock())
    deadline = time.monotonic() + LOCK_TIMEOUT_SECONDS
    if not thread_lock.acquire(timeout=LOCK_TIMEOUT_SECONDS):
        raise RuntimeError("工作簿正由另一个写入任务使用，请稍后重试")

    handle: BinaryIO | None = None
    file_locked = False
    try:
        sidecar = workbook_path.with_name(f".{workbook_path.name}.lock")
        handle = sidecar.open("a+b")
        handle.seek(0, os.SEEK_END)
        if handle.tell() == 0:
            handle.write(b"\0")
            handle.flush()
        while not (file_locked := _try_lock_file(handle)):
            remaining = deadline - time.monotonic()
            if remaining <= 0:
                raise RuntimeError("工作簿正由另一个写入任务使用，请稍后重试")
            time.sleep(min(_LOCK_POLL_SECONDS, remaining))
        yield
    finally:
        try:
            if file_locked and handle is not None:
                _unlock_file(handle)
        finally:
            try:
                if handle is not None:
                    handle.close()
            finally:
                # Unlinking a locked inode can let another process bypass the sidecar.
                thread_lock.release()


class ExcelApplicationRepository:
    def __init__(self, workbook_path: Path, backup_dir: Path) -> None:
        self.workbook_path = Path(workbook_path)
        self.backup_dir = Path(backup_dir)

    def list_applications(self) -> list[TrackerApplication]:
        workbook = self._load_source()
        try:
            sheet = self._sheet(workbook)
            records = self._located_applications(sheet)
        finally:
            workbook.close()
        return [
            located.record
            for located in sorted(
                records,
                key=lambda item: item.record.applied_date or date.min,
                reverse=True,
            )
        ]

    def get_application(self, application_id: str) -> TrackerApplication:
        workbook = self._load_source()
        try:
            sheet = self._sheet(workbook)
            located = self._find_application(sheet, application_id)
            return located.record
        finally:
            workbook.close()

    def create_application(self, draft: TrackerApplicationDraft) -> TrackerApplication:
        record = TrackerApplication(
            id=_application_id(draft.company, draft.role, draft.applied_date),
            **draft.model_dump(mode="python"),
        )

        def mutate(sheet: Worksheet) -> None:
            existing = self._located_applications(sheet)
            if any(item.record.id == record.id for item in existing):
                raise ValueError("重复投递记录：企业、岗位和投递日期已存在")
            self._insert_application(sheet, record, existing)
            if _is_closed(record.status):
                created = self._find_application(sheet, record.id)
                self._move_application_row(
                    sheet,
                    created,
                    record,
                    {},
                    below_divider=True,
                )

        with _exclusive_workbook_lock(self.workbook_path):
            self._mutate(mutate)
            return self.get_application(record.id)

    def update_application(
        self,
        application_id: str,
        patch: TrackerApplicationPatch,
    ) -> TrackerApplication:
        changes = patch.model_dump(mode="python", exclude_unset=True)
        updated_id: str | None = None

        def mutate(sheet: Worksheet) -> None:
            nonlocal updated_id
            existing = self._located_applications(sheet)
            located = next(
                (item for item in existing if item.record.id == application_id),
                None,
            )
            if located is None:
                raise KeyError(application_id)
            values = located.record.model_dump(mode="python")
            values.update(changes)
            values.pop("id")
            prospective_id = _application_id(
                values["company"],
                values["role"],
                values["applied_date"],
            )
            prospective = TrackerApplication(id=prospective_id, **values)
            if any(
                item.row != located.row and item.record.id == prospective_id
                for item in existing
            ):
                raise ValueError("重复投递记录：企业、岗位和投递日期已存在")

            was_closed = _is_closed(located.record.status)
            is_closed = _is_closed(prospective.status)
            if not was_closed and is_closed:
                self._move_application_row(
                    sheet,
                    located,
                    prospective,
                    changes,
                    below_divider=True,
                )
            elif was_closed and not is_closed:
                self._move_application_row(
                    sheet,
                    located,
                    prospective,
                    changes,
                    below_divider=False,
                )
            else:
                self._update_in_place(sheet, located, prospective, changes)
                if is_closed:
                    self._clear_strike(sheet, located.row)
            self._update_company_total(sheet)
            updated_id = prospective_id

        with _exclusive_workbook_lock(self.workbook_path):
            self._mutate(mutate)
            assert updated_id is not None
            return self.get_application(updated_id)

    def _load_source(self):
        try:
            return load_workbook(self.workbook_path)
        except PermissionError as exc:
            raise RuntimeError("无法读取工作簿，请关闭 WPS 中打开的文件后重试") from exc

    @staticmethod
    def _sheet(workbook) -> Worksheet:
        sheet = workbook["投递总览"] if "投递总览" in workbook.sheetnames else workbook.active
        actual = tuple(sheet.cell(1, column).value for column in range(1, len(HEADERS) + 1))
        if actual != HEADERS:
            raise ValueError(f"unexpected headers: {actual!r}; expected {HEADERS!r}")
        extra_headers = [
            sheet.cell(1, column).value
            for column in range(len(HEADERS) + 1, sheet.max_column + 1)
            if sheet.cell(1, column).value not in (None, "")
        ]
        if extra_headers:
            raise ValueError(f"unexpected headers after column I: {extra_headers!r}")
        return sheet

    @staticmethod
    def _is_statistics_row(sheet: Worksheet, row: int) -> bool:
        return str(sheet.cell(row, 1).value or "").strip().startswith("投递公司总数")

    @staticmethod
    def _is_empty_row(sheet: Worksheet, row: int) -> bool:
        return all(
            sheet.cell(row, column).value in (None, "")
            for column in range(1, len(HEADERS) + 1)
        )

    @classmethod
    def _is_divider_row(cls, sheet: Worksheet, row: int) -> bool:
        if not cls._is_empty_row(sheet, row):
            return False
        return any(
            cls._fill_color(sheet.cell(row, column)) == DIVIDER_COLOR
            for column in range(1, len(HEADERS) + 1)
        )

    @staticmethod
    def _fill_color(cell: Cell | MergedCell) -> str | None:
        color = cell.fill.fgColor
        if color.type == "rgb" and color.rgb:
            return color.rgb[-6:].upper()
        return None

    @staticmethod
    def _merged_anchor(sheet: Worksheet, row: int, column: int) -> Cell:
        cell = sheet.cell(row, column)
        if not isinstance(cell, MergedCell):
            return cell
        for merged in sheet.merged_cells.ranges:
            if (
                merged.min_row <= row <= merged.max_row
                and merged.min_col <= column <= merged.max_col
            ):
                return sheet.cell(merged.min_row, merged.min_col)
        raise RuntimeError(f"cannot resolve merged cell at row {row}, column {column}")

    @classmethod
    def _located_applications(cls, sheet: Worksheet) -> list[_LocatedApplication]:
        records = []
        for row in range(2, sheet.max_row + 1):
            if cls._is_statistics_row(sheet, row) or cls._is_empty_row(sheet, row):
                continue
            company_cell = cls._merged_anchor(sheet, row, 1)
            link_cell = cls._merged_anchor(sheet, row, 8)
            company = str(company_cell.value or "").strip()
            role = str(sheet.cell(row, 2).value or "").strip()
            if not company or not role:
                continue
            applied_date = _as_date(sheet.cell(row, 3).value)
            link_target = link_cell.hyperlink.target if link_cell.hyperlink else None
            if link_target is None:
                visible_link = str(link_cell.value or "").strip()
                if visible_link.startswith(("http://", "https://")):
                    link_target = visible_link
            record = TrackerApplication(
                id=_application_id(company, role, applied_date),
                company=company,
                role=role,
                applied_date=applied_date,
                location=cls._optional_text(sheet.cell(row, 4).value),
                status=str(sheet.cell(row, 5).value or "").strip(),
                next_step=cls._optional_text(sheet.cell(row, 6).value),
                next_time=_as_date(sheet.cell(row, 7).value),
                job_url=link_target,
                notes=cls._optional_text(sheet.cell(row, 9).value),
            )
            records.append(_LocatedApplication(row=row, record=record))
        return records

    @staticmethod
    def _optional_text(value: object) -> str | None:
        normalized = str(value or "").strip()
        return normalized or None

    @classmethod
    def _find_application(cls, sheet: Worksheet, application_id: str) -> _LocatedApplication:
        located = next(
            (
                item
                for item in cls._located_applications(sheet)
                if item.record.id == application_id
            ),
            None,
        )
        if located is None:
            raise KeyError(application_id)
        return located

    def _mutate(self, mutation: Callable[[Worksheet], None]) -> None:
        workbook = self._load_source()
        temporary: Path | None = None
        try:
            sheet = self._sheet(workbook)
            mutation(sheet)
            self.backup_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
            backup = self.backup_dir / (
                f"{self.workbook_path.stem}-{timestamp}{self.workbook_path.suffix}"
            )
            shutil.copy2(self.workbook_path, backup)
            temporary = self.workbook_path.with_name(
                f".{self.workbook_path.stem}.{timestamp}{self.workbook_path.suffix}"
            )
            workbook.save(temporary)
            verification = load_workbook(temporary)
            try:
                self._sheet(verification)
            finally:
                verification.close()
            os.replace(temporary, self.workbook_path)
        except PermissionError as exc:
            raise RuntimeError("无法更新工作簿，请关闭 WPS 中打开的文件后重试") from exc
        finally:
            workbook.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    @classmethod
    def _insert_application(
        cls,
        sheet: Worksheet,
        record: TrackerApplication,
        existing: list[_LocatedApplication],
    ) -> None:
        original_max_row = sheet.max_row
        merge_ranges = [copy.copy(merged) for merged in sheet.merged_cells.ranges]
        for merged in list(sheet.merged_cells.ranges):
            if merged.min_row >= 2:
                sheet.unmerge_cells(str(merged))

        dimensions = {
            row: cls._row_dimension(sheet, row)
            for row in range(2, original_max_row + 1)
        }
        first_active_row = next(
            (item.row for item in existing if not _is_closed(item.record.status)),
            existing[0].row if existing else None,
        )
        sheet.insert_rows(2)
        for source_row in range(original_max_row, 1, -1):
            cls._apply_row_dimension(sheet, source_row + 1, dimensions[source_row])
        cls._normalize_hyperlink_refs(sheet)

        style_source_row = first_active_row + 1 if first_active_row is not None else None
        if style_source_row is not None:
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(2, column)._style = copy.copy(
                    sheet.cell(style_source_row, column)._style
                )
            cls._apply_row_dimension(sheet, 2, dimensions[first_active_row])
        cls._write_full_record(sheet, 2, record)

        for merged in merge_ranges:
            if merged.min_row >= 2:
                merged.shift(row_shift=1)
            sheet.merge_cells(str(merged))
        cls._expand_filter(sheet, inserted_at=2)
        cls._update_company_total(sheet)

    @staticmethod
    def _normalize_hyperlink_refs(sheet: Worksheet) -> None:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.hyperlink:
                    cell.hyperlink.ref = cell.coordinate

    @staticmethod
    def _row_dimension(sheet: Worksheet, row: int) -> tuple[float | None, bool, int, bool]:
        dimension = sheet.row_dimensions[row]
        return (
            dimension.height,
            bool(dimension.hidden),
            dimension.outlineLevel,
            bool(dimension.collapsed),
        )

    @staticmethod
    def _apply_row_dimension(
        sheet: Worksheet,
        row: int,
        snapshot: tuple[float | None, bool, int, bool],
    ) -> None:
        dimension = sheet.row_dimensions[row]
        dimension.height, dimension.hidden, dimension.outlineLevel, dimension.collapsed = snapshot

    @staticmethod
    def _expand_filter(sheet: Worksheet, inserted_at: int) -> None:
        if not sheet.auto_filter.ref:
            return
        min_col, min_row, max_col, max_row = range_boundaries(sheet.auto_filter.ref)
        if min_row <= inserted_at <= max_row:
            max_row += 1
        elif inserted_at < min_row:
            min_row += 1
            max_row += 1
        sheet.auto_filter.ref = (
            f"{get_column_letter(min_col)}{min_row}:"
            f"{get_column_letter(max_col)}{max_row}"
        )

    @classmethod
    def _update_in_place(
        cls,
        sheet: Worksheet,
        located: _LocatedApplication,
        prospective: TrackerApplication,
        changes: dict[str, object],
    ) -> None:
        row = located.row
        if {"company", "job_url"}.intersection(changes):
            cls._detach_from_vertical_merges(sheet, row)
        field_columns = {field: column for column, field in enumerate(FIELDS, start=1)}
        for field in changes:
            cls._write_field(sheet, row, field_columns[field], field, getattr(prospective, field))

    @classmethod
    def _detach_from_vertical_merges(
        cls,
        sheet: Worksheet,
        row: int,
    ) -> None:
        for column in (1, 8):
            matching = next(
                (
                    copy.copy(merged)
                    for merged in sheet.merged_cells.ranges
                    if merged.min_col == merged.max_col == column
                    and merged.min_row <= row <= merged.max_row
                    and merged.max_row > merged.min_row
                ),
                None,
            )
            if matching is None:
                continue
            anchor = _cell_snapshot(sheet.cell(matching.min_row, column))
            sheet.unmerge_cells(str(matching))
            _write_cell_snapshot(sheet.cell(row, column), anchor)
            for start_row, end_row in (
                (matching.min_row, row - 1),
                (row + 1, matching.max_row),
            ):
                if start_row > end_row:
                    continue
                _write_cell_snapshot(sheet.cell(start_row, column), anchor)
                if start_row < end_row:
                    sheet.merge_cells(
                        start_row=start_row,
                        end_row=end_row,
                        start_column=column,
                        end_column=column,
                    )

    @staticmethod
    def _clear_strike(sheet: Worksheet, row: int) -> None:
        for column in range(1, len(HEADERS) + 1):
            font = copy.copy(sheet.cell(row, column).font)
            font.strike = False
            sheet.cell(row, column).font = font

    @classmethod
    def _move_application_row(
        cls,
        sheet: Worksheet,
        located: _LocatedApplication,
        prospective: TrackerApplication,
        changes: dict[str, object],
        *,
        below_divider: bool,
    ) -> None:
        original_max_row = sheet.max_row
        max_column = max(sheet.max_column, len(HEADERS))
        rows = [
            cls._snapshot_row(sheet, row, max_column)
            for row in range(2, original_max_row + 1)
        ]
        target = next(row for row in rows if row.source_row == located.row)
        target_cells = list(target.cells)
        for column in (1, 8):
            if isinstance(sheet.cell(located.row, column), MergedCell):
                target_cells[column - 1] = _cell_snapshot(
                    cls._merged_anchor(sheet, located.row, column)
                )
        target.cells = tuple(target_cells)
        merges = [
            _MergeSnapshot(
                min_row=merged.min_row,
                max_row=merged.max_row,
                min_col=merged.min_col,
                max_col=merged.max_col,
                anchor=_cell_snapshot(sheet.cell(merged.min_row, merged.min_col)),
            )
            for merged in sheet.merged_cells.ranges
            if merged.max_row >= 2
        ]
        for merged in list(sheet.merged_cells.ranges):
            if merged.max_row >= 2:
                sheet.unmerge_cells(str(merged))

        ordered = [row for row in rows if row is not target]
        divider: _RowSnapshot | None = None
        created_divider = False
        if below_divider:
            divider_source_row = next(
                (
                    row
                    for row in range(2, original_max_row + 1)
                    if cls._is_divider_row(sheet, row)
                ),
                None,
            )
            divider = next(
                (row for row in ordered if row.source_row == divider_source_row),
                None,
            )
            created_divider = divider is None
            if divider is None:
                divider = cls._new_divider_snapshot(max_column)
                insertion_index = next(
                    (
                        index
                        for index, row in enumerate(ordered)
                        if row.source_row is not None
                        and (
                            cls._is_statistics_snapshot(row)
                            or _is_closed(
                                str(row.cells[4].value or "")
                                if len(row.cells) >= 5
                                else ""
                            )
                        )
                    ),
                    len(ordered),
                )
                ordered.insert(insertion_index, divider)
            divider_index = ordered.index(divider)
            ordered.insert(divider_index + 1, target)
        else:
            ordered.insert(0, target)

        source_to_target = {
            row.source_row: target_row
            for target_row, row in enumerate(ordered, start=2)
            if row.source_row is not None
        }
        for target_row, snapshot in enumerate(ordered, start=2):
            cls._write_row_snapshot(sheet, target_row, snapshot)
        if created_divider:
            assert divider is not None
            divider_row = ordered.index(divider) + 2
            for column in range(1, len(HEADERS) + 1):
                cell = sheet.cell(divider_row, column)
                cell.fill = PatternFill("solid", fgColor=DIVIDER_COLOR)
                cell.font = Font(size=9)
        destination_row = ordered.index(target) + 2
        field_columns = {field: column for column, field in enumerate(FIELDS, start=1)}
        for field in changes:
            cls._write_field(
                sheet,
                destination_row,
                field_columns[field],
                field,
                getattr(prospective, field),
            )
        cls._clear_strike(sheet, destination_row)

        for merged in merges:
            source_rows = list(range(merged.min_row, merged.max_row + 1))
            if located.row in source_rows and len(source_rows) > 1:
                source_rows.remove(located.row)
            mapped_rows = [source_to_target[row] for row in source_rows]
            if not mapped_rows:
                continue
            mapped_rows.sort()
            target_was_merged = located.row in range(
                merged.min_row,
                merged.max_row + 1,
            )
            if target_was_merged and source_rows:
                _write_cell_snapshot(sheet.cell(mapped_rows[0], merged.min_col), merged.anchor)
            if len(mapped_rows) == merged.max_row - merged.min_row + 1 or len(mapped_rows) > 1:
                if mapped_rows == list(range(mapped_rows[0], mapped_rows[-1] + 1)):
                    sheet.merge_cells(
                        start_row=mapped_rows[0],
                        end_row=mapped_rows[-1],
                        start_column=merged.min_col,
                        end_column=merged.max_col,
                    )

        if len(ordered) > original_max_row - 1:
            cls._expand_filter(sheet, inserted_at=destination_row)

    @classmethod
    def _snapshot_row(cls, sheet: Worksheet, row: int, max_column: int) -> _RowSnapshot:
        dimension = sheet.row_dimensions[row]
        return _RowSnapshot(
            source_row=row,
            cells=tuple(
                _cell_snapshot(sheet.cell(row, column))
                for column in range(1, max_column + 1)
            ),
            height=dimension.height,
            hidden=bool(dimension.hidden),
            outline_level=dimension.outlineLevel,
            collapsed=bool(dimension.collapsed),
        )

    @staticmethod
    def _write_row_snapshot(sheet: Worksheet, row: int, snapshot: _RowSnapshot) -> None:
        for column, cell_snapshot in enumerate(snapshot.cells, start=1):
            _write_cell_snapshot(sheet.cell(row, column), cell_snapshot)
        dimension = sheet.row_dimensions[row]
        dimension.height = snapshot.height
        dimension.hidden = snapshot.hidden
        dimension.outlineLevel = snapshot.outline_level
        dimension.collapsed = snapshot.collapsed

    @staticmethod
    def _new_divider_snapshot(max_column: int) -> _RowSnapshot:
        return _RowSnapshot(
            source_row=None,
            cells=tuple(
                _CellSnapshot(value=None, style=None, hyperlink=None, comment=None)
                for _ in range(max_column)
            ),
            height=None,
            hidden=False,
            outline_level=0,
            collapsed=False,
        )

    @staticmethod
    def _is_statistics_snapshot(snapshot: _RowSnapshot) -> bool:
        return str(snapshot.cells[0].value or "").strip().startswith("投递公司总数")

    @classmethod
    def _write_full_record(
        cls,
        sheet: Worksheet,
        row: int,
        record: TrackerApplication,
    ) -> None:
        for column, field in enumerate(FIELDS, start=1):
            cls._write_field(sheet, row, column, field, getattr(record, field))

    @staticmethod
    def _write_field(
        sheet: Worksheet,
        row: int,
        column: int,
        field: str,
        value: object,
    ) -> None:
        cell = sheet.cell(row, column)
        if field == "job_url":
            if value is None:
                cell.value = None
                cell.hyperlink = None
            else:
                cell.value = "查看岗位"
                cell.hyperlink = str(value)
            return
        cell.value = value

    @classmethod
    def _update_company_total(cls, sheet: Worksheet) -> None:
        companies = {
            _normalize(item.record.company) for item in cls._located_applications(sheet)
        }
        for row in range(2, sheet.max_row + 1):
            if cls._is_statistics_row(sheet, row):
                current = str(sheet.cell(row, 1).value)
                if re.search(r"\d+\s*$", current):
                    sheet.cell(row, 1).value = re.sub(r"\d+\s*$", str(len(companies)), current)
                else:
                    sheet.cell(row, 1).value = f"投递公司总数：{len(companies)}"
                return
