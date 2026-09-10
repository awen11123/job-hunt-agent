from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, PatternFill

from scripts.compact_wps_job_tracker import (
    ALLOWED_NEXT_STEPS,
    compact_file_in_place,
    compact_workbook,
    short_next_step,
)


HEADERS = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)


def build_synthetic_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递总览"
    sheet.append(HEADERS)
    rows = [
        ["示例甲", "Agent 工程师", "2026-08-31", "杭州", "applied", "等待示例甲流程通知；准备一大段不应留在总览中的技术内容。", "", "https://example.com/job-a", "重点记录"],
        ["示例乙", "AI 工程师", "2026-08-30", "上海", "AI面试完成", "等待面试结果并补充题型、用时、作答感受和自评。", "", "岗位入口", ""],
        ["示例丙", "算法工程师", "2026-08-29", "北京", "笔试完成", "等待笔试结果并整理全部笔试题目。", "", "有意义文本", "保留批注"],
        ["示例丁", "平台工程师", "2026-08-28", "深圳", "测评异常", "等待有效链接并继续完成测评。", "", "。。。", ""],
        ["示例集团", "大模型应用工程师", "2026-08-27", "杭州", "applied", "等待筛选", "", "统一官网", "第一志愿"],
        [None, "AI 平台开发工程师", "2026-08-27", "杭州", "已投递", "等待集团流程通知；准备后端和 Agent 知识。", "", None, "第二志愿"],
        ["示例终止", "后端工程师", "2026-08-26", "广州", "简历未通过", "已经结束，不再跟进。", "", "", "结束原因"],
    ]
    for row in rows:
        sheet.append(row)

    sheet["H2"].hyperlink = "https://example.com/job-a"
    sheet["H6"].hyperlink = "https://example.com/group"
    sheet["I4"].comment = Comment("这条批注必须保留", "tester")
    sheet["B6"].fill = PatternFill("solid", fgColor="E4DFEC")
    sheet["B6"].alignment = Alignment(wrap_text=True)
    sheet.merge_cells("A6:A7")
    sheet.merge_cells("H6:H7")
    for row in range(2, 9):
        sheet.row_dimensions[row].height = 100
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I8"
    workbook.save(path)


def protected_values(sheet) -> tuple[tuple[object, ...], ...]:
    columns = (1, 2, 3, 4, 7, 9)
    return tuple(
        tuple(sheet.cell(row, column).value for column in columns)
        for row in range(2, sheet.max_row + 1)
    )


def hyperlink_targets(sheet) -> Counter[str]:
    return Counter(
        cell.hyperlink.target
        for row in sheet.iter_rows()
        for cell in row
        if cell.hyperlink
    )


def comments(sheet) -> Counter[tuple[str, str]]:
    return Counter(
        (cell.comment.author, cell.comment.text)
        for row in sheet.iter_rows()
        for cell in row
        if cell.comment
    )


def test_short_next_step_uses_only_approved_process_phrases() -> None:
    cases = [
        ("已投递", "等待筛选", "等待筛选"),
        ("已投递", "一大段准备清单", "等待流程通知"),
        ("AI面试完成", "旧内容", "等待面试结果"),
        ("笔试完成", "旧内容", "等待笔试结果"),
        ("测评异常", "旧内容", "等待补发链接"),
        ("简历未通过", "旧内容", "结束"),
    ]

    for status, current, expected in cases:
        assert short_next_step(status, current) == expected
        assert expected in ALLOWED_NEXT_STEPS


def test_compact_workbook_preserves_structure_and_protected_content(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "compact.xlsx"
    build_synthetic_workbook(source)
    before = load_workbook(source)
    before_sheet = before["投递总览"]
    expected_protected = protected_values(before_sheet)
    expected_links = hyperlink_targets(before_sheet)
    expected_comments = comments(before_sheet)
    expected_merges = {str(cell_range) for cell_range in before_sheet.merged_cells.ranges}
    before.close()

    result = compact_workbook(source, output)

    workbook = load_workbook(output)
    sheet = workbook["投递总览"]
    assert result.rows == 7
    assert result.normalized_statuses == 2
    assert result.shortened_next_steps == 6
    assert result.relabelled_links == 2
    assert result.cleared_placeholders == 1
    assert sheet["E2"].value == "已投递"
    assert sheet["E6"].value == "已投递"
    assert sheet["F2"].value == "等待流程通知"
    assert sheet["F3"].value == "等待面试结果"
    assert sheet["F4"].value == "等待笔试结果"
    assert sheet["F5"].value == "等待补发链接"
    assert sheet["F6"].value == "等待筛选"
    assert sheet["F8"].value == "结束"
    assert sheet["H2"].value == "查看岗位"
    assert sheet["H2"].hyperlink.target == "https://example.com/job-a"
    assert sheet["H3"].value == "岗位入口"
    assert sheet["H5"].value is None
    assert protected_values(sheet) == expected_protected
    assert hyperlink_targets(sheet) == expected_links
    assert comments(sheet) == expected_comments
    assert {str(cell_range) for cell_range in sheet.merged_cells.ranges} == expected_merges
    assert sheet["B6"].fill.fgColor.rgb.endswith("E4DFEC")
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I8"
    assert sheet.column_dimensions["F"].width == 18
    assert sheet.column_dimensions["H"].width == 12
    assert max(sheet.row_dimensions[row].height for row in range(2, 9)) <= 42
    assert sheet["H2"].alignment.horizontal == "center"
    assert sheet["H2"].alignment.wrap_text is not True
    assert [sheet.cell(row, 1).value for row in range(2, 9)] == [
        "示例甲",
        "示例乙",
        "示例丙",
        "示例丁",
        "示例集团",
        None,
        "示例终止",
    ]
    workbook.close()


def test_compact_file_in_place_creates_exact_backup(tmp_path: Path) -> None:
    target = tmp_path / "tracker.xlsx"
    backup_dir = tmp_path / "backups"
    build_synthetic_workbook(target)
    original = target.read_bytes()

    backup, result = compact_file_in_place(target, backup_dir)

    assert backup.parent == backup_dir
    assert backup.name.startswith("tracker-compact-")
    assert backup.read_bytes() == original
    assert result.rows == 7
    assert load_workbook(target)["投递总览"]["H2"].value == "查看岗位"
