from __future__ import annotations

import copy
import threading
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import ValidationError

import job_hunt_agent.excel.repository as repository_module
from job_hunt_agent.excel.models import (
    TrackerApplicationDraft,
    TrackerApplicationPatch,
)
from job_hunt_agent.excel.repository import ExcelApplicationRepository


HEADERS = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)
DIVIDER_COLOR = "E7E6E3"


def build_synthetic_tracker(path: Path, *, include_divider: bool = True) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递总览"
    sheet.append(HEADERS)
    sheet.append(
        [
            "示例科技",
            "Agent 工程师",
            date(2026, 9, 8),
            "杭州",
            "已投递",
            "等待筛选",
            None,
            "查看岗位",
            "保留备注",
        ]
    )
    sheet.append(
        [
            "示例集团",
            "后端工程师",
            date(2026, 9, 6),
            "上海",
            "笔试",
            "等待笔试结果",
            date(2026, 9, 10),
            "查看岗位",
            "第一志愿",
        ]
    )
    sheet.append(
        [
            None,
            "平台工程师",
            date(2026, 9, 7),
            "北京",
            "面试",
            "准备一面",
            date(2026, 9, 12),
            None,
            "第二志愿",
        ]
    )
    if include_divider:
        sheet.append([None] * len(HEADERS))
        divider_row = sheet.max_row
        for cell in sheet[divider_row]:
            cell.fill = PatternFill("solid", fgColor=DIVIDER_COLOR)
            cell.font = Font(name="微软雅黑", size=9)
    sheet.append(
        [
            "旧公司",
            "旧岗位",
            date(2026, 8, 20),
            "深圳",
            "简历未通过",
            "结束",
            None,
            "查看岗位",
            "历史记录",
        ]
    )
    sheet.append([None] * len(HEADERS))
    sheet.append(["投递公司总数：3", None, None, None, None, None, None, None, None])

    sheet["H2"].hyperlink = "https://example.com/job"
    sheet["H3"].hyperlink = "https://example.com/group"
    closed_row = 6 if include_divider else 5
    sheet.cell(closed_row, 8).hyperlink = "https://example.com/old"
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("H3:H4")

    sheet["A2"].fill = PatternFill("solid", fgColor="C6E0B4")
    sheet["A3"].fill = PatternFill("solid", fgColor="DDEBF7")
    sheet["B2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet["H3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["I2"].comment = Comment("必须保留", "tester")
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 31 + row
    sheet.column_dimensions["B"].width = 36
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"

    metadata = workbook.create_sheet("只读元数据")
    metadata["A1"] = "不要修改"
    workbook.save(path)
    workbook.close()
    return path


def build_five_row_merged_tracker(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递总览"
    sheet.append(HEADERS)
    sheet.append(
        [
            "其他公司",
            "其他岗位",
            date(2026, 9, 9),
            "杭州",
            "已投递",
            "等待筛选",
            None,
            None,
            None,
        ]
    )
    for index in range(5):
        sheet.append(
            [
                " 三行集团 " if index == 0 else None,
                f"合并岗位-{index + 1}",
                date(2026, 9, index + 1),
                "上海",
                "已投递",
                "等待筛选",
                None,
                "集团入口" if index == 0 else None,
                f"备注-{index + 1}",
            ]
        )
    sheet.append([None] * len(HEADERS))
    for cell in sheet[sheet.max_row]:
        cell.fill = PatternFill("solid", fgColor=DIVIDER_COLOR)
        cell.font = Font(size=9)
    sheet.append(
        [
            "旧公司",
            "旧岗位",
            date(2026, 8, 1),
            None,
            "结束",
            "结束",
            None,
            None,
            None,
        ]
    )
    sheet.append(["投递公司总数：3", None, None, None, None, None, None, None, None])
    sheet["A3"].fill = PatternFill("solid", fgColor="DDEBF7")
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["H3"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet["H3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["H3"].hyperlink = "https://example.com/group-five"
    sheet.merge_cells("A3:A7")
    sheet.merge_cells("H3:H7")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I10"
    workbook.save(path)
    workbook.close()
    return path


def record_by_role(repository: ExcelApplicationRepository, role: str):
    return next(record for record in repository.list_applications() if record.role == role)


def synchronize_first_workbook_loads(
    monkeypatch: pytest.MonkeyPatch,
    workbook_path: Path,
) -> None:
    real_load_workbook = repository_module.load_workbook
    barrier = threading.Barrier(2)
    seen_threads: set[int] = set()
    guard = threading.Lock()

    def synchronized_load(path: str | Path, *args, **kwargs):
        workbook = real_load_workbook(path, *args, **kwargs)
        should_wait = False
        if Path(path).resolve() == workbook_path.resolve():
            thread_id = threading.get_ident()
            with guard:
                if thread_id not in seen_threads:
                    seen_threads.add(thread_id)
                    should_wait = True
        if should_wait:
            try:
                barrier.wait(timeout=0.2)
            except threading.BrokenBarrierError:
                pass
        return workbook

    monkeypatch.setattr(repository_module, "load_workbook", synchronized_load)


def test_concurrent_creates_from_two_repositories_do_not_lose_updates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    first = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    second = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    synchronize_first_workbook_loads(monkeypatch, workbook_path)

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [
            executor.submit(
                repository.create_application,
                TrackerApplicationDraft(
                    company=f"并发公司-{index}",
                    role=f"并发岗位-{index}",
                    applied_date=date(2026, 9, 10 + index),
                ),
            )
            for index, repository in enumerate((first, second), start=1)
        ]
        created = [future.result() for future in futures]

    assert {record.role for record in created} == {"并发岗位-1", "并发岗位-2"}
    assert {record.role for record in first.list_applications()}.issuperset(
        {"并发岗位-1", "并发岗位-2"}
    )


def test_concurrent_updates_from_two_repositories_do_not_lose_updates(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    first = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    second = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    first_target = record_by_role(first, "Agent 工程师")
    second_target = record_by_role(second, "后端工程师")
    synchronize_first_workbook_loads(monkeypatch, workbook_path)

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [
            executor.submit(
                first.update_application,
                first_target.id,
                TrackerApplicationPatch(notes="并发更新-1"),
            ),
            executor.submit(
                second.update_application,
                second_target.id,
                TrackerApplicationPatch(notes="并发更新-2"),
            ),
        ]
        updated = [future.result() for future in futures]

    assert {record.notes for record in updated} == {"并发更新-1", "并发更新-2"}
    records = first.list_applications()
    assert next(record for record in records if record.id == first_target.id).notes == "并发更新-1"
    assert next(record for record in records if record.id == second_target.id).notes == "并发更新-2"


def test_write_lock_timeout_has_clear_private_error(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    monkeypatch.setattr(repository_module, "LOCK_TIMEOUT_SECONDS", 0.01, raising=False)
    monkeypatch.setattr(
        repository_module,
        "_try_lock_file",
        lambda _handle: False,
        raising=False,
    )

    with pytest.raises(RuntimeError, match="另一个写入任务") as error:
        repository.create_application(
            TrackerApplicationDraft(company="锁测试", role="锁测试岗位")
        )

    assert str(workbook_path) not in str(error.value)


def test_unlock_failure_still_releases_handle_and_thread_lock(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    first = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    second = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    real_unlock = repository_module._unlock_file
    unlock_calls = 0

    def fail_first_unlock(handle) -> None:
        nonlocal unlock_calls
        unlock_calls += 1
        real_unlock(handle)
        if unlock_calls == 1:
            raise OSError("simulated unlock failure")

    monkeypatch.setattr(repository_module, "_unlock_file", fail_first_unlock)
    monkeypatch.setattr(repository_module, "LOCK_TIMEOUT_SECONDS", 0.05)

    with pytest.raises(OSError, match="simulated unlock failure"):
        first.create_application(
            TrackerApplicationDraft(company="首次写入", role="首次岗位")
        )

    created = second.create_application(
        TrackerApplicationDraft(company="后续写入", role="后续岗位")
    )

    assert created.role == "后续岗位"
    assert workbook_path.with_name(f".{workbook_path.name}.lock").read_bytes() == b"\0"


def test_repository_lists_newest_first_and_inherits_merged_company_and_link(
    tmp_path: Path,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")

    records = repository.list_applications()

    assert [record.role for record in records] == [
        "Agent 工程师",
        "平台工程师",
        "后端工程师",
        "旧岗位",
    ]
    platform = records[1]
    assert platform.company == "示例集团"
    assert str(platform.job_url) == "https://example.com/group"
    assert records[0].notes == "保留备注"


def test_stable_id_does_not_depend_on_row_number(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    original = record_by_role(repository, "Agent 工程师")

    repository.create_application(
        TrackerApplicationDraft(
            company="新公司",
            role="新岗位",
            applied_date=date(2026, 9, 9),
        )
    )

    shifted = record_by_role(repository, "Agent 工程师")
    assert shifted.id == original.id
    assert shifted.id.startswith("app_")
    assert len(shifted.id) == 20


def test_create_rejects_normalized_duplicate_without_writing(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    original = workbook_path.read_bytes()

    with pytest.raises(ValueError, match="重复"):
        repository.create_application(
            TrackerApplicationDraft(
                company="  示例科技  ",
                role="AGENT   工程师",
                applied_date=date(2026, 9, 8),
            )
        )

    assert workbook_path.read_bytes() == original
    assert not (tmp_path / "backups").exists()


def test_create_makes_exact_backup_and_preserves_layout(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    backup_dir = tmp_path / "backups"
    original_bytes = workbook_path.read_bytes()
    before = load_workbook(workbook_path)
    before_sheet = before["投递总览"]
    expected_style = before_sheet["A2"]._style
    expected_row_height = before_sheet.row_dimensions[2].height
    before.close()
    repository = ExcelApplicationRepository(workbook_path, backup_dir)

    created = repository.create_application(
        TrackerApplicationDraft(
            company="示例旅行",
            role="Agent 平台工程师",
            applied_date=date(2026, 9, 9),
            job_url="https://example.com/new",
        )
    )

    backups = list(backup_dir.glob("*.xlsx"))
    assert len(backups) == 1
    assert backups[0].read_bytes() == original_bytes
    assert created == repository.get_application(created.id)

    after = load_workbook(workbook_path)
    sheet = after["投递总览"]
    assert sheet["A2"].value == "示例旅行"
    assert sheet["H2"].value == "查看岗位"
    assert sheet["H2"].hyperlink.target == "https://example.com/new"
    assert sheet["A2"]._style == expected_style
    assert sheet.row_dimensions[2].height == expected_row_height
    assert sheet["A3"].value == "示例科技"
    assert sheet["I3"].comment.text == "必须保留"
    assert {str(cell_range) for cell_range in sheet.merged_cells.ranges} == {
        "A4:A5",
        "H4:H5",
    }
    assert sheet["H4"].hyperlink.target == "https://example.com/group"
    assert sheet["A6"].value is None
    assert sheet["A6"].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    assert sheet["A9"].value == "投递公司总数：4"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I9"
    assert sheet.column_dimensions["B"].width == 36
    assert after["只读元数据"]["A1"].value == "不要修改"
    after.close()


def test_create_closed_application_places_it_first_below_existing_divider(
    tmp_path: Path,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    before = load_workbook(workbook_path)
    before_sheet = before["投递总览"]
    source_font = copy.copy(before_sheet["A2"].font)
    source_font.strike = False
    before_sheet["A2"].font = source_font
    before.save(workbook_path)
    expected_style = copy.copy(before_sheet["A2"]._style)
    before.close()
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")

    created = repository.create_application(
        TrackerApplicationDraft(
            company="新结束公司",
            role="新结束岗位",
            applied_date=date(2026, 9, 10),
            status="简历未通过",
            next_step="结束",
            job_url="https://example.com/new-closed",
        )
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert sheet["B2"].value == "Agent 工程师"
    assert sheet["A5"].value is None
    assert sheet["A5"].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    assert sheet["B6"].value == "新结束岗位"
    assert sheet["B7"].value == "旧岗位"
    assert sheet["A6"]._style == expected_style
    assert sheet["H6"].value == "查看岗位"
    assert sheet["H6"].hyperlink.target == "https://example.com/new-closed"
    assert not any(sheet.cell(6, column).font.strike for column in range(1, 10))
    assert sheet["A9"].value == "投递公司总数：4"
    assert sum(
        1
        for row in range(2, sheet.max_row + 1)
        if sheet["A" + str(row)].value is None
        and sheet["A" + str(row)].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    ) == 1
    assert created == repository.get_application(created.id)
    workbook.close()


def test_update_changes_only_requested_fields_in_place(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "Agent 工程师")
    before = load_workbook(workbook_path)
    expected_style = before["投递总览"]["B2"]._style
    before.close()

    updated = repository.update_application(
        target.id,
        TrackerApplicationPatch(
            location="南京",
            next_step="等待面试",
            job_url="https://example.com/updated",
        ),
    )

    assert updated.id == target.id
    assert updated.location == "南京"
    assert updated.next_step == "等待面试"
    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert sheet["A2"].value == "示例科技"
    assert sheet["B2"]._style == expected_style
    assert sheet["I2"].value == "保留备注"
    assert sheet["I2"].comment.text == "必须保留"
    assert sheet["H2"].value == "查看岗位"
    assert sheet["H2"].hyperlink.target == "https://example.com/updated"
    workbook.close()


def test_reopening_closed_application_moves_it_to_top_and_preserves_row(
    tmp_path: Path,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    sheet["A6"].fill = PatternFill("solid", fgColor="F4CCCC")
    company_font = copy.copy(sheet["A6"].font)
    company_font.strike = False
    sheet["A6"].font = company_font
    sheet["H6"].value = "历史入口"
    sheet["H6"].alignment = Alignment(horizontal="center", vertical="center")
    link_font = copy.copy(sheet["H6"].font)
    link_font.strike = False
    sheet["H6"].font = link_font
    sheet["H6"].comment = Comment("历史链接批注", "tester")
    workbook.save(workbook_path)
    workbook.close()

    before = load_workbook(workbook_path)
    expected_company_style = copy.copy(before["投递总览"]["A6"]._style)
    expected_link_style = copy.copy(before["投递总览"]["H6"]._style)
    before.close()
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "旧岗位")

    reopened = repository.update_application(
        target.id,
        TrackerApplicationPatch(status="重新开放"),
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert reopened.status == "重新开放"
    assert sheet["B2"].value == "旧岗位"
    assert sheet["E2"].value == "重新开放"
    assert sheet["A2"]._style == expected_company_style
    assert sheet["H2"].value == "历史入口"
    assert sheet["H2"].hyperlink.target == "https://example.com/old"
    assert sheet["H2"]._style == expected_link_style
    assert sheet["H2"].comment.text == "历史链接批注"
    assert sheet["A6"].value is None
    assert sheet["A6"].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    assert sheet["A8"].value == "投递公司总数：3"
    assert sum(
        1
        for row in range(2, sheet.max_row + 1)
        if sheet["A" + str(row)].value is None
        and sheet["A" + str(row)].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    ) == 1
    workbook.close()


def test_closed_to_closed_update_stays_in_its_existing_row(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    sheet.insert_rows(7)
    sheet["A7"] = "更早公司"
    sheet["B7"] = "更早岗位"
    sheet["C7"] = date(2026, 8, 1)
    sheet["E7"] = "结束"
    sheet["F7"] = "结束"
    sheet.auto_filter.ref = "A1:I9"
    workbook.save(workbook_path)
    workbook.close()
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "更早岗位")

    repository.update_application(
        target.id,
        TrackerApplicationPatch(status="流程终止", notes="原因更新"),
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert sheet["B6"].value == "旧岗位"
    assert sheet["B7"].value == "更早岗位"
    assert sheet["E7"].value == "流程终止"
    assert sheet["I7"].value == "原因更新"
    workbook.close()


@pytest.mark.parametrize(
    ("changes", "expected_company", "expected_label", "expected_target"),
    [
        (
            {"company": "新中间公司"},
            "新中间公司",
            "集团入口",
            "https://example.com/group-five",
        ),
        (
            {"job_url": "https://example.com/new-middle"},
            " 三行集团 ",
            "查看岗位",
            "https://example.com/new-middle",
        ),
    ],
)
def test_updating_middle_of_long_merge_splits_siblings_on_each_side(
    tmp_path: Path,
    changes: dict[str, str],
    expected_company: str,
    expected_label: str,
    expected_target: str,
) -> None:
    workbook_path = build_five_row_merged_tracker(tmp_path / "tracker.xlsx")
    before = load_workbook(workbook_path)
    before_sheet = before["投递总览"]
    expected_company_style = copy.copy(before_sheet["A3"]._style)
    expected_link_style = copy.copy(before_sheet["H3"]._style)
    before.close()
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "合并岗位-3")

    repository.update_application(target.id, TrackerApplicationPatch(**changes))

    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert {str(cell_range) for cell_range in sheet.merged_cells.ranges} == {
        "A3:A4",
        "A6:A7",
        "H3:H4",
        "H6:H7",
    }
    assert sheet["A3"].value == " 三行集团 "
    assert sheet["A6"].value == " 三行集团 "
    assert sheet["A5"].value == expected_company
    assert sheet["A3"]._style == expected_company_style
    assert sheet["A6"]._style == expected_company_style
    assert sheet["A5"]._style == expected_company_style
    assert sheet["H3"].value == "集团入口"
    assert sheet["H6"].value == "集团入口"
    assert sheet["H5"].value == expected_label
    assert sheet["H3"].hyperlink.target == "https://example.com/group-five"
    assert sheet["H6"].hyperlink.target == "https://example.com/group-five"
    assert sheet["H5"].hyperlink.target == expected_target
    assert sheet["H3"]._style == expected_link_style
    assert sheet["H6"]._style == expected_link_style
    assert sheet["H5"]._style == expected_link_style
    assert [sheet[f"B{row}"].value for row in range(3, 8)] == [
        "合并岗位-1",
        "合并岗位-2",
        "合并岗位-3",
        "合并岗位-4",
        "合并岗位-5",
    ]
    workbook.close()


def test_closing_one_merged_role_splits_it_and_moves_it_below_existing_divider(
    tmp_path: Path,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "平台工程师")
    before = load_workbook(workbook_path)
    expected_company_fill = before["投递总览"]["A3"].fill.fgColor.rgb
    expected_link_alignment = before["投递总览"]["H3"].alignment.horizontal
    before.close()

    updated = repository.update_application(
        target.id,
        TrackerApplicationPatch(status="流程结束", next_step="结束"),
    )

    assert updated.status == "流程结束"
    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert sheet["B3"].value == "后端工程师"
    assert sheet["A3"].value == "示例集团"
    assert sheet["H3"].hyperlink.target == "https://example.com/group"
    assert sheet["A4"].value is None
    assert sheet["A4"].fill.fgColor.rgb.endswith(DIVIDER_COLOR)
    assert sheet["B5"].value == "平台工程师"
    assert sheet["A5"].value == "示例集团"
    assert sheet["H5"].value == "查看岗位"
    assert sheet["H5"].hyperlink.target == "https://example.com/group"
    assert sheet["A5"].fill.fgColor.rgb == expected_company_fill
    assert sheet["H5"].alignment.horizontal == expected_link_alignment
    assert sheet["B6"].value == "旧岗位"
    assert not any(
        sheet.cell(5, column).font.strike for column in range(1, len(HEADERS) + 1)
    )
    assert not sheet.merged_cells.ranges
    assert sheet["A8"].value == "投递公司总数：3"
    workbook.close()


def test_closing_preserves_unpatched_raw_cell_content_and_formatting(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    raw_values = (
        " 甲 ",
        " 原岗位 ",
        "2026-09-08",
        " 杭州 ",
        "已投递",
        " 等待筛选 ",
        "2026-09-10",
        "自定义入口",
        "  保留备注  ",
    )
    for column, value in enumerate(raw_values, start=1):
        cell = sheet.cell(2, column)
        cell.value = value
        cell.fill = PatternFill("solid", fgColor=f"{column:02X}{column:02X}{column:02X}")
        cell.font = Font(name="Arial", size=10 + column, strike=False)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.comment = Comment(f"批注-{column}", "tester")
    sheet["H2"].hyperlink = "https://example.com/raw-target"
    sheet.row_dimensions[2].height = 47
    workbook.save(workbook_path)
    workbook.close()

    before = load_workbook(workbook_path)
    before_sheet = before["投递总览"]
    expected = {
        column: (
            before_sheet.cell(2, column).value,
            type(before_sheet.cell(2, column).value),
            copy.copy(before_sheet.cell(2, column)._style),
            before_sheet.cell(2, column).hyperlink.target
            if before_sheet.cell(2, column).hyperlink
            else None,
            before_sheet.cell(2, column).comment.text,
            before_sheet.cell(2, column).comment.author,
        )
        for column in range(1, 10)
        if column != 5
    }
    expected_status_style = copy.copy(before_sheet["E2"]._style)
    before.close()

    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "原岗位")
    repository.update_application(
        target.id,
        TrackerApplicationPatch(status="流程结束"),
    )

    after = load_workbook(workbook_path)
    moved = after["投递总览"]
    assert moved["E5"].value == "流程结束"
    assert moved["E5"]._style == expected_status_style
    assert moved["E5"].comment.text == "批注-5"
    assert moved["E5"].comment.author == "tester"
    assert moved.row_dimensions[5].height == 47
    for column, signature in expected.items():
        cell = moved.cell(5, column)
        assert (
            cell.value,
            type(cell.value),
            cell._style,
            cell.hyperlink.target if cell.hyperlink else None,
            cell.comment.text,
            cell.comment.author,
        ) == signature
    assert all(moved.cell(5, column).font.strike is False for column in range(1, 10))
    after.close()


def test_first_closed_update_creates_gray_divider(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(
        tmp_path / "tracker.xlsx",
        include_divider=False,
    )
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    target = record_by_role(repository, "Agent 工程师")

    repository.update_application(
        target.id,
        TrackerApplicationPatch(status="主动放弃"),
    )

    workbook = load_workbook(workbook_path)
    sheet = workbook["投递总览"]
    assert sheet["B3"].value == "平台工程师"
    assert all(sheet.cell(4, column).value is None for column in range(1, 10))
    assert all(
        sheet.cell(4, column).fill.fgColor.rgb.endswith(DIVIDER_COLOR)
        for column in range(1, 10)
    )
    assert all(sheet.cell(4, column).font.sz == 9 for column in range(1, 10))
    assert sheet["B5"].value == "Agent 工程师"
    assert sheet["H2"].hyperlink.target == "https://example.com/group"
    assert sheet["A8"].value == "投递公司总数：3"
    workbook.close()


def test_models_reject_empty_patch_and_unknown_fields() -> None:
    assert TrackerApplicationDraft(company="甲", role="乙").applied_date == date.today()
    assert TrackerApplicationDraft(company="甲", role="乙").status == "已投递"
    assert TrackerApplicationDraft(company="甲", role="乙").next_step == "等待筛选"

    with pytest.raises(ValidationError, match="empty patch"):
        TrackerApplicationPatch()
    with pytest.raises(ValidationError, match="extra_forbidden"):
        TrackerApplicationPatch(unknown="value")


@pytest.mark.parametrize("field_name", ("company", "role", "status"))
def test_patch_rejects_explicit_null_for_required_text_fields(field_name: str) -> None:
    with pytest.raises(ValidationError, match=f"{field_name} cannot be null"):
        TrackerApplicationPatch(**{field_name: None})


def test_patch_allows_explicit_null_for_clearable_fields() -> None:
    patch = TrackerApplicationPatch(location=None, next_step=None, job_url=None)

    assert patch.model_fields_set == {"location", "next_step", "job_url"}


def test_unknown_id_raises_key_error(tmp_path: Path) -> None:
    repository = ExcelApplicationRepository(
        build_synthetic_tracker(tmp_path / "tracker.xlsx"),
        tmp_path / "backups",
    )

    with pytest.raises(KeyError, match="missing"):
        repository.get_application("missing")
    with pytest.raises(KeyError, match="missing"):
        repository.update_application(
            "missing",
            TrackerApplicationPatch(status="面试"),
        )


def test_bad_headers_raise_clear_value_error(tmp_path: Path) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    workbook = load_workbook(workbook_path)
    workbook["投递总览"]["B1"] = "错误列名"
    workbook.save(workbook_path)
    workbook.close()

    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")

    with pytest.raises(ValueError, match="unexpected headers"):
        repository.list_applications()


def test_wps_file_lock_keeps_original_and_cleans_temporary_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    original = workbook_path.read_bytes()

    def fail_replace(source: str | Path, destination: str | Path) -> None:
        raise PermissionError("simulated WPS lock")

    monkeypatch.setattr(repository_module.os, "replace", fail_replace)

    with pytest.raises(RuntimeError, match="关闭 WPS"):
        repository.create_application(
            TrackerApplicationDraft(company="新公司", role="新岗位")
        )

    assert workbook_path.read_bytes() == original
    assert not list(tmp_path.glob(".tracker.*.xlsx"))


def test_validation_failure_keeps_original_and_cleans_temporary_file(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook_path = build_synthetic_tracker(tmp_path / "tracker.xlsx")
    repository = ExcelApplicationRepository(workbook_path, tmp_path / "backups")
    original = workbook_path.read_bytes()
    real_load_workbook = repository_module.load_workbook

    def fail_temporary_validation(path: str | Path, *args, **kwargs):
        if Path(path).name.startswith(".tracker."):
            raise OSError("simulated validation failure")
        return real_load_workbook(path, *args, **kwargs)

    monkeypatch.setattr(repository_module, "load_workbook", fail_temporary_validation)

    with pytest.raises(OSError, match="simulated validation failure"):
        repository.update_application(
            record_by_role(repository, "Agent 工程师").id,
            TrackerApplicationPatch(notes="new notes"),
        )

    assert workbook_path.read_bytes() == original
    assert not list(tmp_path.glob(".tracker.*.xlsx"))
