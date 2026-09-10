from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from scripts.optimize_wps_job_tracker import (
    optimize_file_in_place,
    optimize_workbook,
)


HEADERS = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)


def build_synthetic_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递总览"
    sheet.append(HEADERS)
    rows = [
        ["示例出行", "Agent 工程师", "2026-08-30", "杭州", "已投递", "等待结果", "", "岗位页面", "保留原备注"],
        ["示例终止甲", "算法工程师", "2026-08-27", "上海", "简历未通过", "结束", "", "结束岗位", ""],
        ["示例平台", "AI 工程师", "2026-08-29", "深圳", "测评异常", "等待补发链接", "", "平台岗位", "不要沉底"],
        ["示例联通", "大模型应用", "2026-08-28", "杭州", "applied", "等待筛选", "", "统一官网", "第一志愿"],
        [None, "AI 平台开发", "2026-08-28", "杭州", "applied", "等待筛选", "", None, "第二志愿"],
        ["示例终止乙", "后端工程师", "2026-08-26", "北京", "结束", "结束", "", "另一岗位", ""],
    ]
    for row in rows:
        sheet.append(row)

    links = {
        "H2": "https://example.com/travel",
        "H3": "https://example.com/closed-a",
        "H4": "https://example.com/platform",
        "H5": "https://example.com/group",
        "H7": "https://example.com/closed-b",
    }
    for coordinate, target in links.items():
        sheet[coordinate].hyperlink = target
        sheet[coordinate].font = Font(color="0563C1", underline="single")

    sheet["I3"].comment = Comment("结束原因需要保留", "tester")
    sheet["B5"].fill = PatternFill("solid", fgColor="E4DFEC")
    sheet["B5"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[5].height = 42
    sheet.row_dimensions[6].height = 28
    sheet.merge_cells("A5:A6")
    sheet.merge_cells("H5:H6")
    sheet.freeze_panes = "A10"
    sheet.auto_filter.ref = "A1:I7"
    workbook.save(path)


def row_content_signature(sheet, row: int) -> tuple[tuple[object, ...], tuple[str | None, ...]]:
    values = tuple(sheet.cell(row, column).value for column in range(1, 10))
    links = tuple(
        sheet.cell(row, column).hyperlink.target
        if sheet.cell(row, column).hyperlink
        else None
        for column in range(1, 10)
    )
    return values, links


def normalized_signatures(path: Path, *, updated: bool) -> Counter:
    sheet = load_workbook(path)["投递总览"]
    signatures = []
    for row in range(2, sheet.max_row + 1):
        values, links = row_content_signature(sheet, row)
        values = list(values)
        if updated and values[0] == "示例出行":
            values[4] = "AI面试完成"
        signatures.append((tuple(values), links))
    return Counter(signatures)


def test_optimize_workbook_preserves_content_links_and_group_merges(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "optimized.xlsx"
    build_synthetic_workbook(source)
    expected = normalized_signatures(source, updated=True)

    result = optimize_workbook(
        source,
        output,
        status_updates={"示例出行": "AI面试完成"},
    )

    sheet = load_workbook(output)["投递总览"]
    assert result.active_rows == 4
    assert result.closed_rows == 2
    assert result.updated_companies == ("示例出行",)
    assert [sheet[f"A{row}"].value for row in (2, 3, 4, 6, 7)] == [
        "示例出行",
        "示例平台",
        "示例联通",
        "示例终止甲",
        "示例终止乙",
    ]
    assert sheet["E2"].value == "AI面试完成"
    assert sheet["F2"].value == "等待结果"
    assert sheet["H2"].value == "岗位页面"
    assert sheet["H2"].hyperlink.target == "https://example.com/travel"
    assert sheet["I6"].comment.text == "结束原因需要保留"
    assert {str(cell_range) for cell_range in sheet.merged_cells.ranges} == {
        "A4:A5",
        "H4:H5",
    }
    assert sheet.row_dimensions[4].height == 42
    assert sheet.row_dimensions[5].height == 28
    assert sheet["B4"].fill.fgColor.rgb.endswith("E4DFEC")
    assert normalized_signatures(output, updated=False) == expected


def test_optimize_workbook_applies_stage_colors_and_closed_boundary(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "optimized.xlsx"
    build_synthetic_workbook(source)

    optimize_workbook(
        source,
        output,
        status_updates={"示例出行": "AI面试完成"},
    )

    sheet = load_workbook(output)["投递总览"]
    colors = {sheet[f"E{row}"].fill.fgColor.rgb for row in (2, 3, 4, 6)}
    assert len(colors) == 4
    assert sheet["E2"].fill.fgColor.rgb.endswith("C6EFCE")
    assert sheet["E3"].fill.fgColor.rgb.endswith("FCE4D6")
    assert sheet["E4"].fill.fgColor.rgb.endswith("DDEBF7")
    assert sheet["E6"].fill.fgColor.rgb.endswith("F4CCCC")
    assert sheet["F2"].fill.fgColor.rgb != sheet["E2"].fill.fgColor.rgb
    assert all(sheet.cell(6, column).border.top.style == "thick" for column in range(1, 10))
    assert sheet["A6"].fill.fgColor.rgb.endswith("F2F4F7")
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I7"


def test_optimize_file_in_place_creates_recoverable_backup(tmp_path: Path) -> None:
    target = tmp_path / "tracker.xlsx"
    backup_dir = tmp_path / "backups"
    build_synthetic_workbook(target)
    original_bytes = target.read_bytes()

    backup, result = optimize_file_in_place(
        target,
        backup_dir,
        status_updates={"示例出行": "AI面试完成"},
    )

    assert backup.parent == backup_dir
    assert backup.name.startswith("tracker-")
    assert backup.read_bytes() == original_bytes
    assert result.active_rows == 4
    assert load_workbook(target)["投递总览"]["E2"].value == "AI面试完成"
