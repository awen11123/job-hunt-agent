import os
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.wps_job_tracker import (
    HEADERS,
    build_workbook,
    create_sync_probe,
    discover_wps_account_directory,
    main,
)


def sample_record(*, company: str, applied_date: str) -> dict[str, str]:
    return {
        "company": company,
        "role": "Agent 工程师",
        "applied_date": applied_date,
        "location": "杭州",
        "status": "已投递",
        "next_step": "等待通知",
        "next_time": "",
        "link": "",
        "notes": "",
    }


def test_build_workbook_layout_sorting_and_links(tmp_path: Path) -> None:
    output = tmp_path / "tracker.xlsx"
    records = [
        {
            "company": "较早企业",
            "role": "AI 工程师",
            "applied_date": "2026-08-01",
            "location": "北京",
            "status": "已投递",
            "next_step": "等待通知",
            "next_time": "",
            "link": "https://example.com/old",
            "notes": "",
        },
        {
            "company": "最新企业",
            "role": "Agent 工程师",
            "applied_date": "2026-08-30",
            "location": "杭州",
            "status": "面试安排中",
            "next_step": "确认时间",
            "next_time": "2026-09-01 10:00",
            "link": "https://example.com/new",
            "notes": "重点跟进",
        },
    ]

    build_workbook(records, output)

    workbook = load_workbook(output)
    sheet = workbook["投递总览"]
    assert [cell.value for cell in sheet[1]] == list(HEADERS)
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:I3"
    assert sheet["A2"].value == "最新企业"
    assert sheet["A3"].value == "较早企业"
    assert sheet["H2"].hyperlink.target == "https://example.com/new"
    assert sheet["E2"].fill.fgColor.rgb != sheet["E3"].fill.fgColor.rgb


def test_build_workbook_merges_only_adjacent_identical_companies(
    tmp_path: Path,
) -> None:
    output = tmp_path / "tracker.xlsx"
    records = [
        sample_record(company="同一企业", applied_date="2026-08-30"),
        sample_record(company="其他企业", applied_date="2026-08-30"),
        sample_record(company="同一企业", applied_date="2026-08-30"),
        sample_record(company="同一企业", applied_date="2026-08-28"),
    ]

    build_workbook(records, output)

    sheet = load_workbook(output)["投递总览"]
    assert {str(cell_range) for cell_range in sheet.merged_cells.ranges} == {"A2:A3"}
    assert sheet["A2"].value == "同一企业"
    assert sheet["A4"].value == "其他企业"
    assert sheet["A5"].value == "同一企业"


def test_failed_save_preserves_existing_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "tracker.xlsx"
    build_workbook(
        [sample_record(company="第一版", applied_date="2026-08-29")], output
    )

    def interrupted_save(_workbook: Workbook, filename: str | Path) -> None:
        Path(filename).write_bytes(b"incomplete workbook")
        raise OSError("simulated interrupted save")

    monkeypatch.setattr(Workbook, "save", interrupted_save)
    with pytest.raises(OSError, match="simulated interrupted save"):
        build_workbook(
            [sample_record(company="第二版", applied_date="2026-08-30")], output
        )
    monkeypatch.undo()

    workbook = load_workbook(output)
    assert workbook["投递总览"]["A2"].value == "第一版"


def test_build_workbook_rejects_unexpected_fields(tmp_path: Path) -> None:
    record = sample_record(company="示例企业", applied_date="2026-08-30")
    record["unexpected"] = "不应写入"

    with pytest.raises(ValueError, match="exactly the nine normalized fields"):
        build_workbook([record], tmp_path / "tracker.xlsx")


def test_locked_workbook_reports_how_to_retry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "tracker.xlsx"
    build_workbook(
        [sample_record(company="第一版", applied_date="2026-08-29")], output
    )

    def deny_replace(_source: str | Path, _target: str | Path) -> None:
        raise PermissionError("file is locked")

    monkeypatch.setattr(os, "replace", deny_replace)
    with pytest.raises(RuntimeError, match="关闭 WPS"):
        build_workbook(
            [sample_record(company="第二版", applied_date="2026-08-30")], output
        )


def test_discover_wps_account_directory_ignores_cache_directories(
    tmp_path: Path,
) -> None:
    (tmp_path / ".metadata").mkdir()
    (tmp_path / "hyperionlocalcache").mkdir()
    account = tmp_path / "account-directory"
    account.mkdir()

    assert discover_wps_account_directory(tmp_path) == account


def test_discover_wps_account_directory_rejects_ambiguity(tmp_path: Path) -> None:
    (tmp_path / "account-one").mkdir()
    (tmp_path / "account-two").mkdir()

    with pytest.raises(RuntimeError, match="multiple WPS account directories"):
        discover_wps_account_directory(tmp_path)


def test_discover_wps_account_directory_reports_missing_account(
    tmp_path: Path,
) -> None:
    (tmp_path / ".metadata").mkdir()
    (tmp_path / "hyperionlocalcache").mkdir()

    with pytest.raises(RuntimeError, match="no WPS account directory"):
        discover_wps_account_directory(tmp_path)


def test_create_sync_probe_writes_only_fake_data(tmp_path: Path) -> None:
    account = tmp_path / "account-directory"
    account.mkdir()

    output = create_sync_probe("同步测试-第一版", root=tmp_path)

    assert output == account / "WPS同步测试.xlsx"
    workbook = load_workbook(output)
    sheet = workbook["投递总览"]
    assert sheet.max_row == 2
    assert sheet["A2"].value == "同步测试"
    assert sheet["I2"].value == "同步测试-第一版"


def test_probe_command_uses_logged_in_cloud_directory(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    home = tmp_path / "home"
    account = home / "WPS Cloud Files" / "account-directory"
    account.mkdir(parents=True)
    monkeypatch.setattr(Path, "home", lambda: home)

    assert main(["probe", "--marker", "命令行测试"]) == 0

    output = account / "WPS同步测试.xlsx"
    assert output.exists()
    assert str(output) in capsys.readouterr().out


def test_build_command_reads_normalized_json(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    source = tmp_path / "applications.json"
    output = tmp_path / "tracker.xlsx"
    source.write_text(
        json.dumps(
            [sample_record(company="示例企业", applied_date="2026-08-30")],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    assert main(["build", "--input", str(source), "--output", str(output)]) == 0

    workbook = load_workbook(output)
    assert workbook["投递总览"]["A2"].value == "示例企业"
    report = capsys.readouterr().out
    assert str(output) in report
    assert "1 records" in report


def test_probe_command_can_open_saved_workbook(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "home"
    account = home / "WPS Cloud Files" / "account-directory"
    account.mkdir(parents=True)
    monkeypatch.setattr(Path, "home", lambda: home)
    opened: list[Path] = []
    monkeypatch.setattr(os, "startfile", lambda path: opened.append(Path(path)))

    assert main(["probe", "--marker", "打开测试", "--open"]) == 0

    assert opened == [account / "WPS同步测试.xlsx"]
