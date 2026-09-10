from __future__ import annotations

import argparse
import copy
import os
import shutil
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Final, Mapping

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.comments import Comment
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet


HEADERS: Final = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)
DEFAULT_STATUS_UPDATES: Final = {
    "示例科技": "AI面试完成",
    "示例旅行": "AI面试完成",
}
CLOSED_MARKERS: Final = ("未通过", "结束", "拒绝", "放弃", "终止", "淘汰")

_BLUE = PatternFill("solid", fgColor="DDEBF7")
_BLUE_LIGHT = PatternFill("solid", fgColor="EDF4FB")
_YELLOW = PatternFill("solid", fgColor="FFF2CC")
_YELLOW_LIGHT = PatternFill("solid", fgColor="FFF9E6")
_GREEN = PatternFill("solid", fgColor="C6EFCE")
_GREEN_LIGHT = PatternFill("solid", fgColor="EAF7ED")
_ORANGE = PatternFill("solid", fgColor="FCE4D6")
_ORANGE_LIGHT = PatternFill("solid", fgColor="FEF2EB")
_RED = PatternFill("solid", fgColor="F4CCCC")
_RED_LIGHT = PatternFill("solid", fgColor="FCE8E8")
_MUTED = PatternFill("solid", fgColor="F2F4F7")
_THICK_BOUNDARY = Side(style="thick", color="667085")


@dataclass(frozen=True)
class OptimizationResult:
    active_rows: int
    closed_rows: int
    updated_companies: tuple[str, ...]


@dataclass
class CellSnapshot:
    value: object
    style: object
    hyperlink: object
    comment: Comment | None


@dataclass
class RowSnapshot:
    cells: tuple[CellSnapshot, ...]
    height: float | None
    hidden: bool
    outline_level: int
    collapsed: bool


@dataclass
class MergeSnapshot:
    min_row_offset: int
    max_row_offset: int
    min_col: int
    max_col: int


@dataclass
class RowGroup:
    source_start: int
    rows: tuple[RowSnapshot, ...]
    merges: tuple[MergeSnapshot, ...]
    company: str
    closed: bool


def _is_closed(status: object) -> bool:
    normalized = str(status or "").strip()
    return bool(normalized) and any(marker in normalized for marker in CLOSED_MARKERS)


def _cell_snapshot(cell: Cell | MergedCell) -> CellSnapshot:
    return CellSnapshot(
        value=cell.value,
        style=copy.copy(cell._style),
        hyperlink=copy.copy(getattr(cell, "hyperlink", None)),
        comment=copy.copy(getattr(cell, "comment", None)),
    )


def _row_groups(sheet: Worksheet) -> list[RowGroup]:
    first_row = 2
    last_row = sheet.max_row
    parents = {row: row for row in range(first_row, last_row + 1)}

    def find(row: int) -> int:
        while parents[row] != row:
            parents[row] = parents[parents[row]]
            row = parents[row]
        return row

    def union(left: int, right: int) -> None:
        left_root = find(left)
        right_root = find(right)
        if left_root != right_root:
            parents[right_root] = left_root

    data_merges = []
    for merged_range in sheet.merged_cells.ranges:
        if merged_range.max_row < first_row or merged_range.min_row > last_row:
            continue
        data_merges.append(copy.copy(merged_range))
        start = max(first_row, merged_range.min_row)
        end = min(last_row, merged_range.max_row)
        for row in range(start + 1, end + 1):
            union(start, row)

    components: dict[int, list[int]] = {}
    for row in range(first_row, last_row + 1):
        components.setdefault(find(row), []).append(row)

    groups = []
    for row_numbers in sorted(components.values(), key=lambda rows: rows[0]):
        if row_numbers != list(range(row_numbers[0], row_numbers[-1] + 1)):
            raise ValueError("merged row groups must be contiguous")
        start = row_numbers[0]
        row_snapshots = []
        for row in row_numbers:
            dimension = sheet.row_dimensions[row]
            row_snapshots.append(
                RowSnapshot(
                    cells=tuple(
                        _cell_snapshot(sheet.cell(row, column))
                        for column in range(1, len(HEADERS) + 1)
                    ),
                    height=dimension.height,
                    hidden=bool(dimension.hidden),
                    outline_level=dimension.outlineLevel,
                    collapsed=bool(dimension.collapsed),
                )
            )
        merges = tuple(
            MergeSnapshot(
                min_row_offset=merged.min_row - start,
                max_row_offset=merged.max_row - start,
                min_col=merged.min_col,
                max_col=merged.max_col,
            )
            for merged in data_merges
            if merged.min_row >= start and merged.max_row <= row_numbers[-1]
        )
        company = next(
            (
                str(snapshot.cells[0].value).strip()
                for snapshot in row_snapshots
                if snapshot.cells[0].value not in (None, "")
            ),
            "",
        )
        statuses = [snapshot.cells[4].value for snapshot in row_snapshots]
        groups.append(
            RowGroup(
                source_start=start,
                rows=tuple(row_snapshots),
                merges=merges,
                company=company,
                closed=bool(statuses) and all(_is_closed(status) for status in statuses),
            )
        )
    return groups


def _signature_from_snapshot(
    group: RowGroup,
    row: RowSnapshot,
    status_updates: Mapping[str, str],
) -> tuple[tuple[object, str | None, str | None, str | None], ...]:
    signature = []
    for column, snapshot in enumerate(row.cells, start=1):
        value = snapshot.value
        if column == 5 and group.company in status_updates:
            value = status_updates[group.company]
        hyperlink = getattr(snapshot.hyperlink, "target", None)
        comment_text = snapshot.comment.text if snapshot.comment else None
        comment_author = snapshot.comment.author if snapshot.comment else None
        signature.append((value, hyperlink, comment_text, comment_author))
    return tuple(signature)


def _signature_from_sheet(sheet: Worksheet, row: int) -> tuple:
    signature = []
    for column in range(1, len(HEADERS) + 1):
        cell = sheet.cell(row, column)
        signature.append(
            (
                cell.value,
                cell.hyperlink.target if cell.hyperlink else None,
                cell.comment.text if cell.comment else None,
                cell.comment.author if cell.comment else None,
            )
        )
    return tuple(signature)


def _write_group(sheet: Worksheet, group: RowGroup, target_start: int) -> None:
    for offset, row_snapshot in enumerate(group.rows):
        target_row = target_start + offset
        dimension = sheet.row_dimensions[target_row]
        dimension.height = row_snapshot.height
        dimension.hidden = row_snapshot.hidden
        dimension.outlineLevel = row_snapshot.outline_level
        dimension.collapsed = row_snapshot.collapsed
        for column, snapshot in enumerate(row_snapshot.cells, start=1):
            cell = sheet.cell(target_row, column)
            cell.value = snapshot.value
            cell._style = copy.copy(snapshot.style)
            cell.hyperlink = copy.copy(snapshot.hyperlink)
            cell.comment = copy.copy(snapshot.comment)


def _stage_fills(status: object) -> tuple[PatternFill, PatternFill]:
    normalized = str(status or "").strip()
    if _is_closed(normalized):
        return _RED, _RED_LIGHT
    if any(marker in normalized for marker in ("异常", "待处理", "补发")):
        return _ORANGE, _ORANGE_LIGHT
    if "AI面试完成" in normalized or "面试完成" in normalized:
        return _GREEN, _GREEN_LIGHT
    if any(marker in normalized for marker in ("笔试", "测评")):
        return _YELLOW, _YELLOW_LIGHT
    if normalized in {"已投递", "applied"} or "投递" in normalized:
        return _BLUE, _BLUE_LIGHT
    return _BLUE, _BLUE_LIGHT


def _replace_top_border(cell: Cell, top: Side) -> None:
    border = cell.border
    cell.border = Border(
        left=border.left,
        right=border.right,
        top=top,
        bottom=border.bottom,
        diagonal=border.diagonal,
        diagonal_direction=border.diagonal_direction,
        diagonalUp=border.diagonalUp,
        diagonalDown=border.diagonalDown,
        outline=border.outline,
        vertical=border.vertical,
        horizontal=border.horizontal,
    )


def _apply_visuals(sheet: Worksheet, first_closed_row: int | None) -> None:
    for row in range(2, sheet.max_row + 1):
        status_cell = sheet.cell(row, 5)
        primary_fill, secondary_fill = _stage_fills(status_cell.value)
        if _is_closed(status_cell.value):
            for column in range(1, len(HEADERS) + 1):
                sheet.cell(row, column).fill = copy.copy(_MUTED)
        status_cell.fill = copy.copy(primary_fill)
        status_cell.font = copy.copy(status_cell.font)
        status_cell.font = Font(
            name=status_cell.font.name,
            size=status_cell.font.size,
            bold=True,
            italic=status_cell.font.italic,
            vertAlign=status_cell.font.vertAlign,
            underline=status_cell.font.underline,
            strike=status_cell.font.strike,
            color=status_cell.font.color,
        )
        for column in (6, 7):
            sheet.cell(row, column).fill = copy.copy(secondary_fill)

    if first_closed_row is not None:
        for column in range(1, len(HEADERS) + 1):
            _replace_top_border(sheet.cell(first_closed_row, column), _THICK_BOUNDARY)


def _validate_headers(sheet: Worksheet) -> None:
    actual = tuple(sheet.cell(1, column).value for column in range(1, 10))
    if actual != HEADERS:
        raise ValueError(f"unexpected headers: {actual!r}")


def optimize_workbook(
    source: Path,
    output: Path,
    *,
    status_updates: Mapping[str, str] | None = None,
) -> OptimizationResult:
    source = Path(source)
    output = Path(output)
    updates = dict(DEFAULT_STATUS_UPDATES if status_updates is None else status_updates)
    workbook = load_workbook(source)
    sheet = workbook["投递总览"] if "投递总览" in workbook.sheetnames else workbook.active
    _validate_headers(sheet)
    groups = _row_groups(sheet)
    active_groups = [group for group in groups if not group.closed]
    closed_groups = [group for group in groups if group.closed]
    ordered_groups = active_groups + closed_groups
    expected_signatures = Counter(
        _signature_from_snapshot(group, row, updates)
        for group in groups
        for row in group.rows
    )

    data_merges = [
        str(merged)
        for merged in sheet.merged_cells.ranges
        if merged.max_row >= 2
    ]
    for merged in data_merges:
        sheet.unmerge_cells(merged)

    translated_merges = []
    target_row = 2
    updated_companies = []
    for group in ordered_groups:
        _write_group(sheet, group, target_row)
        if group.company in updates:
            updated_companies.append(group.company)
            for offset in range(len(group.rows)):
                sheet.cell(target_row + offset, 5).value = updates[group.company]
        for merged in group.merges:
            translated_merges.append(
                (
                    target_row + merged.min_row_offset,
                    target_row + merged.max_row_offset,
                    merged.min_col,
                    merged.max_col,
                )
            )
        target_row += len(group.rows)

    actual_signatures = Counter(
        _signature_from_sheet(sheet, row) for row in range(2, sheet.max_row + 1)
    )
    if actual_signatures != expected_signatures:
        raise RuntimeError("content or hyperlinks changed outside the authorized status updates")

    active_rows = sum(len(group.rows) for group in active_groups)
    closed_rows = sum(len(group.rows) for group in closed_groups)
    first_closed_row = 2 + active_rows if closed_rows else None
    _apply_visuals(sheet, first_closed_row)
    for min_row, max_row, min_col, max_col in translated_merges:
        sheet.merge_cells(
            start_row=min_row,
            end_row=max_row,
            start_column=min_col,
            end_column=max_col,
        )

    sheet.freeze_panes = "A2"
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    verification = load_workbook(output, read_only=True)
    verification.close()
    return OptimizationResult(
        active_rows=active_rows,
        closed_rows=closed_rows,
        updated_companies=tuple(dict.fromkeys(updated_companies)),
    )


def optimize_file_in_place(
    target: Path,
    backup_dir: Path,
    *,
    status_updates: Mapping[str, str] | None = None,
) -> tuple[Path, OptimizationResult]:
    target = Path(target)
    backup_dir = Path(backup_dir)
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup = backup_dir / f"{target.stem}-{timestamp}{target.suffix}"
    shutil.copy2(target, backup)
    temporary = target.with_name(f".{target.stem}.optimize-{timestamp}{target.suffix}")
    try:
        result = optimize_workbook(
            target,
            temporary,
            status_updates=status_updates,
        )
        os.replace(temporary, target)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise RuntimeError("无法替换工作簿，请关闭 WPS 中打开的文件后重试") from exc
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return backup, result


def inspect_workbook(target: Path) -> OptimizationResult:
    workbook = load_workbook(target, read_only=False)
    sheet = workbook["投递总览"] if "投递总览" in workbook.sheetnames else workbook.active
    _validate_headers(sheet)
    groups = _row_groups(sheet)
    workbook.close()
    active_rows = sum(len(group.rows) for group in groups if not group.closed)
    closed_rows = sum(len(group.rows) for group in groups if group.closed)
    present_updates = tuple(
        company for company in DEFAULT_STATUS_UPDATES if any(group.company == company for group in groups)
    )
    return OptimizationResult(active_rows, closed_rows, present_updates)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Optimize the private WPS job tracker in place")
    parser.add_argument("--target", required=True, type=Path)
    parser.add_argument("--backup-dir", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    if args.dry_run:
        result = inspect_workbook(args.target)
        print(
            f"dry-run: active_rows={result.active_rows}, closed_rows={result.closed_rows}, "
            f"status_updates={','.join(result.updated_companies)}"
        )
        return 0

    if args.backup_dir is None:
        parser.error("--backup-dir is required unless --dry-run is used")
    backup, result = optimize_file_in_place(args.target, args.backup_dir)
    print(
        f"optimized: active_rows={result.active_rows}, closed_rows={result.closed_rows}, "
        f"status_updates={','.join(result.updated_companies)}, backup={backup}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
