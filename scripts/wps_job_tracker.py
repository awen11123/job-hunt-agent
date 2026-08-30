from __future__ import annotations

import argparse
import json
import os
from datetime import date
from pathlib import Path
from typing import Final

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


HEADERS: Final = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)

FIELDS: Final = (
    "company",
    "role",
    "applied_date",
    "location",
    "status",
    "next_step",
    "next_time",
    "link",
    "notes",
)

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_ACTIVE_FILL = PatternFill("solid", fgColor="C6E0B4")
_ATTENTION_FILL = PatternFill("solid", fgColor="FFE699")
_CLOSED_FILL = PatternFill("solid", fgColor="F4CCCC")
_THIN_GRAY = Side(style="thin", color="D9E2F3")


def _status_fill(status: str) -> PatternFill:
    if any(word in status for word in ("未通过", "结束", "放弃", "拒绝")):
        return _CLOSED_FILL
    if any(word in status for word in ("测评", "笔试", "面试", "待处理")):
        return _ATTENTION_FILL
    return _ACTIVE_FILL


def discover_wps_account_directory(root: Path | None = None) -> Path:
    root = root or Path.home() / "WPS Cloud Files"
    candidates = [
        path
        for path in root.iterdir()
        if path.is_dir()
        and not path.name.startswith(".")
        and path.name.lower() != "hyperionlocalcache"
    ]
    if len(candidates) > 1:
        raise RuntimeError("multiple WPS account directories found")
    if not candidates:
        raise RuntimeError("no WPS account directory found")
    return candidates[0]


def create_sync_probe(marker: str, root: Path | None = None) -> Path:
    account_directory = discover_wps_account_directory(root)
    record = {
        "company": "同步测试",
        "role": "WPS 固定链接验证",
        "applied_date": date.today().isoformat(),
        "location": "本地",
        "status": "待处理",
        "next_step": "检查只读链接",
        "next_time": "",
        "link": "",
        "notes": marker,
    }
    return build_workbook([record], account_directory / "WPS同步测试.xlsx")


def build_workbook(records: list[dict[str, str]], output: Path) -> Path:
    """Render records newest-first and atomically replace output."""
    expected_fields = set(FIELDS)
    for record in records:
        if set(record) != expected_fields:
            raise ValueError("each record must contain exactly the nine normalized fields")

    ordered = sorted(records, key=lambda record: record["applied_date"], reverse=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "投递总览"
    sheet.append(HEADERS)

    for record in ordered:
        sheet.append([record[field] for field in FIELDS])
        row = sheet.max_row
        sheet.cell(row, 5).fill = _status_fill(record["status"])
        if record["link"]:
            link_cell = sheet.cell(row, 8)
            link_cell.hyperlink = record["link"]
            link_cell.style = "Hyperlink"

    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = Border(bottom=_THIN_GRAY)
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    widths = (20, 34, 13, 16, 14, 24, 19, 24, 30)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
    sheet.row_dimensions[1].height = 24
    sheet.sheet_view.showGridLines = False

    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.name}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, output)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise RuntimeError(
            "无法更新工作簿；请关闭 WPS 中打开的文件后重试"
        ) from exc
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return output


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the private WPS job tracker")
    commands = parser.add_subparsers(dest="command", required=True)
    probe = commands.add_parser("probe", help="create a fake sync probe")
    probe.add_argument("--marker", required=True)
    probe.add_argument("--open", action="store_true")
    build = commands.add_parser("build", help="build the private job tracker")
    build.add_argument("--input", required=True, type=Path)
    build.add_argument("--output", type=Path)
    args = parser.parse_args(argv)

    if args.command == "probe":
        output = create_sync_probe(args.marker)
        print(output)
        if args.open:
            os.startfile(output)
        return 0

    if args.command == "build":
        records = json.loads(args.input.read_text(encoding="utf-8"))
        if not isinstance(records, list):
            raise ValueError("input JSON must contain a list of records")
        output = args.output or (
            discover_wps_account_directory() / "秋招投递总览.xlsx"
        )
        build_workbook(records, output)
        print(f"{output} ({len(records)} records)")
        return 0

    raise AssertionError(f"unsupported command: {args.command}")


if __name__ == "__main__":
    raise SystemExit(main())
