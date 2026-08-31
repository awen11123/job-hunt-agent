from __future__ import annotations

import argparse
import copy
import os
import shutil
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Final

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.worksheet.worksheet import Worksheet


HEADERS: Final = (
    "企业",
    "投递岗位",
    "投递日期",
    "所在地",
    "当前状态",
    "下一节点",
    "节点时间",
    "岗位链接",
    "备注",
)
ALLOWED_NEXT_STEPS: Final = frozenset(
    {
        "等待筛选",
        "等待流程通知",
        "等待笔试结果",
        "等待面试结果",
        "等待补发链接",
        "结束",
    }
)
CLOSED_MARKERS: Final = ("未通过", "结束", "拒绝", "放弃", "终止", "淘汰")
INVALID_LINK_PLACEHOLDERS: Final = frozenset({"。。。"})
PROTECTED_COLUMNS: Final = (1, 2, 3, 4, 7, 9)


@dataclass(frozen=True)
class CompactResult:
    rows: int
    normalized_statuses: int
    shortened_next_steps: int
    relabelled_links: int
    cleared_placeholders: int


@dataclass(frozen=True)
class WorkbookInvariants:
    rows: int
    columns: int
    protected_values: tuple[tuple[object, ...], ...]
    hyperlink_targets: Counter[str]
    comments: Counter[tuple[str, str]]
    meaningful_nonlink_values: tuple[tuple[int, object], ...]
    merged_ranges: frozenset[str]
    freeze_panes: str | None
    filter_ref: str | None
    sheet_names: tuple[str, ...]


def _is_closed(status: object) -> bool:
    normalized = str(status or "").strip()
    return bool(normalized) and any(marker in normalized for marker in CLOSED_MARKERS)


def short_next_step(status: object, current: object) -> str:
    normalized = str(status or "").strip()
    if _is_closed(normalized):
        return "结束"
    if "AI面试完成" in normalized or "面试完成" in normalized:
        return "等待面试结果"
    if "笔试完成" in normalized:
        return "等待笔试结果"
    if any(marker in normalized for marker in ("异常", "补发")):
        return "等待补发链接"
    if str(current or "").strip() == "等待筛选":
        return "等待筛选"
    return "等待流程通知"


def _sheet(workbook) -> Worksheet:
    return workbook["投递总览"] if "投递总览" in workbook.sheetnames else workbook.active


def _validate_headers(sheet: Worksheet) -> None:
    actual = tuple(sheet.cell(1, column).value for column in range(1, 10))
    if actual != HEADERS:
        raise ValueError(f"unexpected headers: {actual!r}")


def _hyperlink_targets(sheet: Worksheet) -> Counter[str]:
    return Counter(
        cell.hyperlink.target
        for row in sheet.iter_rows()
        for cell in row
        if getattr(cell, "hyperlink", None)
    )


def _comments(sheet: Worksheet) -> Counter[tuple[str, str]]:
    return Counter(
        (cell.comment.author, cell.comment.text)
        for row in sheet.iter_rows()
        for cell in row
        if getattr(cell, "comment", None)
    )


def _snapshot(workbook, sheet: Worksheet) -> WorkbookInvariants:
    protected_values = tuple(
        tuple(sheet.cell(row, column).value for column in PROTECTED_COLUMNS)
        for row in range(2, sheet.max_row + 1)
    )
    meaningful_nonlink_values = []
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, 8)
        if isinstance(cell, MergedCell) or cell.hyperlink:
            continue
        value = cell.value
        if value is not None and str(value).strip() not in INVALID_LINK_PLACEHOLDERS:
            meaningful_nonlink_values.append((row, value))
    return WorkbookInvariants(
        rows=sheet.max_row,
        columns=sheet.max_column,
        protected_values=protected_values,
        hyperlink_targets=_hyperlink_targets(sheet),
        comments=_comments(sheet),
        meaningful_nonlink_values=tuple(meaningful_nonlink_values),
        merged_ranges=frozenset(str(cell_range) for cell_range in sheet.merged_cells.ranges),
        freeze_panes=str(sheet.freeze_panes) if sheet.freeze_panes else None,
        filter_ref=sheet.auto_filter.ref,
        sheet_names=tuple(workbook.sheetnames),
    )


def _validate_invariants(
    workbook,
    sheet: Worksheet,
    expected: WorkbookInvariants,
) -> None:
    actual = _snapshot(workbook, sheet)
    if actual != expected:
        raise RuntimeError("protected workbook content or structure changed")
    next_steps = {
        str(sheet.cell(row, 6).value or "").strip()
        for row in range(2, sheet.max_row + 1)
    }
    if not next_steps.issubset(ALLOWED_NEXT_STEPS):
        raise RuntimeError("an unapproved next-step phrase remains")
    heights = [
        sheet.row_dimensions[row].height
        for row in range(2, sheet.max_row + 1)
        if sheet.row_dimensions[row].height is not None
    ]
    if heights and max(heights) > 42:
        raise RuntimeError("a data row is taller than the compact layout limit")


def _planned_result(sheet: Worksheet) -> CompactResult:
    normalized_statuses = 0
    shortened_next_steps = 0
    relabelled_links = 0
    cleared_placeholders = 0
    for row in range(2, sheet.max_row + 1):
        status_cell = sheet.cell(row, 5)
        if str(status_cell.value or "").strip() == "applied":
            normalized_statuses += 1
        next_step = short_next_step(status_cell.value, sheet.cell(row, 6).value)
        if sheet.cell(row, 6).value != next_step:
            shortened_next_steps += 1
        link_cell = sheet.cell(row, 8)
        if isinstance(link_cell, MergedCell):
            continue
        if link_cell.hyperlink and link_cell.value != "查看岗位":
            relabelled_links += 1
        elif (
            not link_cell.hyperlink
            and str(link_cell.value or "").strip() in INVALID_LINK_PLACEHOLDERS
        ):
            cleared_placeholders += 1
    return CompactResult(
        rows=sheet.max_row - 1,
        normalized_statuses=normalized_statuses,
        shortened_next_steps=shortened_next_steps,
        relabelled_links=relabelled_links,
        cleared_placeholders=cleared_placeholders,
    )


def _compact_sheet(sheet: Worksheet) -> CompactResult:
    result = _planned_result(sheet)
    for row in range(2, sheet.max_row + 1):
        status_cell = sheet.cell(row, 5)
        if str(status_cell.value or "").strip() == "applied":
            status_cell.value = "已投递"

        next_cell = sheet.cell(row, 6)
        next_cell.value = short_next_step(status_cell.value, next_cell.value)

        link_cell = sheet.cell(row, 8)
        if not isinstance(link_cell, MergedCell):
            if link_cell.hyperlink:
                link_cell.value = "查看岗位"
                alignment = copy.copy(link_cell.alignment)
                alignment.horizontal = "center"
                alignment.vertical = "center"
                alignment.wrap_text = False
                link_cell.alignment = alignment
            elif str(link_cell.value or "").strip() in INVALID_LINK_PLACEHOLDERS:
                link_cell.value = None

        company = str(sheet.cell(row, 1).value or "")
        role = str(sheet.cell(row, 2).value or "")
        if _is_closed(status_cell.value):
            height = 28
        elif len(company) > 14 or len(role) > 18:
            height = 42
        else:
            height = 32
        sheet.row_dimensions[row].height = height

    sheet.column_dimensions["F"].width = 18
    sheet.column_dimensions["H"].width = 12
    return result


def compact_workbook(source: Path, output: Path) -> CompactResult:
    source = Path(source)
    output = Path(output)
    if source.resolve() == output.resolve():
        raise ValueError("source and output must be different paths")

    workbook = load_workbook(source)
    try:
        sheet = _sheet(workbook)
        _validate_headers(sheet)
        expected = _snapshot(workbook, sheet)
        result = _compact_sheet(sheet)
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(output)
    finally:
        workbook.close()

    verification = load_workbook(output)
    try:
        verification_sheet = _sheet(verification)
        _validate_headers(verification_sheet)
        _validate_invariants(verification, verification_sheet, expected)
    except BaseException:
        output.unlink(missing_ok=True)
        raise
    finally:
        verification.close()
    return result


def compact_file_in_place(
    target: Path,
    backup_dir: Path,
) -> tuple[Path, CompactResult]:
    target = Path(target)
    backup_dir = Path(backup_dir)
    backup_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    backup = backup_dir / f"{target.stem}-compact-{timestamp}{target.suffix}"
    shutil.copy2(target, backup)
    temporary = target.with_name(f".{target.stem}.compact-{timestamp}{target.suffix}")
    try:
        result = compact_workbook(target, temporary)
        os.replace(temporary, target)
    except PermissionError as exc:
        temporary.unlink(missing_ok=True)
        raise RuntimeError("无法替换工作簿，请关闭 WPS 中打开的文件后重试") from exc
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    return backup, result


def inspect_plan(target: Path) -> CompactResult:
    workbook = load_workbook(target)
    try:
        sheet = _sheet(workbook)
        _validate_headers(sheet)
        return _planned_result(sheet)
    finally:
        workbook.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Compact the private WPS job tracker")
    parser.add_argument("--target", required=True, type=Path)
    parser.add_argument("--backup-dir", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)

    if args.dry_run:
        result = inspect_plan(args.target)
        print(
            f"dry-run: rows={result.rows}, normalized_statuses={result.normalized_statuses}, "
            f"shortened_next_steps={result.shortened_next_steps}, "
            f"relabelled_links={result.relabelled_links}, "
            f"cleared_placeholders={result.cleared_placeholders}"
        )
        return 0
    if args.backup_dir is None:
        parser.error("--backup-dir is required unless --dry-run is used")

    backup, result = compact_file_in_place(args.target, args.backup_dir)
    print(
        f"compacted: rows={result.rows}, normalized_statuses={result.normalized_statuses}, "
        f"shortened_next_steps={result.shortened_next_steps}, "
        f"relabelled_links={result.relabelled_links}, "
        f"cleared_placeholders={result.cleared_placeholders}, backup={backup}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
